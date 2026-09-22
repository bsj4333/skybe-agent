from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader

TEXT_EXT = {".md", ".txt", ".csv", ".py", ".json", ".yaml", ".yml"}
SKIP_DIRS = {".git", "__pycache__", "node_modules", ".venv"}
MAX_FILE_BYTES = 2_000_000
LINES_PER_CHUNK = 80


@dataclass(frozen=True)
class Chunk:
    source: str
    location: str
    text: str


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("cp949", errors="replace")


def _text_chunks(source: str, text: str) -> list[Chunk]:
    lines = text.splitlines()
    chunks = []
    for start in range(0, len(lines), LINES_PER_CHUNK):
        part = lines[start : start + LINES_PER_CHUNK]
        chunks.append(Chunk(source, f"L{start + 1}-{start + len(part)}", "\n".join(part)))
    return chunks


def _pdf_chunks(source: str, path: Path) -> list[Chunk]:
    reader = PdfReader(path)
    return [
        Chunk(source, f"p.{i}", (page.extract_text() or "").strip())
        for i, page in enumerate(reader.pages, start=1)
    ]


def _pptx_chunks(source: str, path: Path) -> list[Chunk]:
    chunks = []
    for i, slide in enumerate(Presentation(path).slides, start=1):
        texts = [
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        chunks.append(Chunk(source, f"slide {i}", "\n".join(texts)))
    return chunks


def load_folder(root: Path) -> list[Chunk]:
    chunks: list[Chunk] = []
    for path in sorted(root.rglob("*")):
        if not path.is_file() or any(part in SKIP_DIRS for part in path.parts):
            continue
        if path.stat().st_size > MAX_FILE_BYTES:
            continue
        source = path.relative_to(root).as_posix()
        suffix = path.suffix.lower()
        if suffix in TEXT_EXT:
            chunks += _text_chunks(source, _read_text(path))
        elif suffix == ".pdf":
            chunks += _pdf_chunks(source, path)
        elif suffix == ".pptx":
            chunks += _pptx_chunks(source, path)
    return [c for c in chunks if c.text.strip()]
