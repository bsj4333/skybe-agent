from pathlib import Path

from fpdf import FPDF
from pptx import Presentation

BASE = Path(__file__).parent / "3d_printer_raw"


def make_pptx() -> None:
    prs = Presentation()
    slides = [
        ("3D 프린터 PHM 이상탐지", ["출력 중 센서 데이터로 불량을 실시간 탐지", "기계공학과 캡스톤디자인"]),
        ("문제", ["출력 종료 후에야 불량 확인", "재료·시간 낭비"]),
        ("결과", ["F1-score 0.71 → 0.86", "추론 지연 평균 42ms", "불량 발견 출력 중 3분 이내"]),
        ("한계와 향후 과제", ["PLA 소재, 0.4mm 노즐 한 가지 조건에서만 검증", "다양한 소재로 확장 필요"]),
    ]
    for title, bullets in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "\n".join(bullets)
    prs.save(BASE / "발표자료.pptx")


def make_pdf() -> None:
    pdf = FPDF()
    pdf.add_font("Malgun", fname="C:/Windows/Fonts/malgun.ttf")
    pdf.add_page()
    pdf.set_font("Malgun", size=12)
    lines = [
        "주간 회의록 (2025.05.14)",
        "",
        "안건 1. 임계값 설정",
        "- 임계값을 검증 데이터 95퍼센타일로 잡았더니 정상 출력에서 오탐이 잦았음",
        "- 99퍼센타일로 올리자 오탐이 줄었고 F1이 올랐음. 김민수가 재실험 후 확정",
        "",
        "안건 2. 데이터 부족",
        "- 불량 데이터가 적어 정상 데이터만으로 학습하는 방식(재구성 오차)을 채택",
        "- 이서연이 불량 출력 실험을 10회 추가 진행하기로 함",
    ]
    for line in lines:
        pdf.multi_cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(BASE / "회의록.pdf"))


if __name__ == "__main__":
    make_pptx()
    make_pdf()
    print("샘플 생성 완료")
